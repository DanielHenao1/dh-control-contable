import io

from django import forms
from django.contrib import messages
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from controles.reglas_exogena import pagos_acumulados, tope
from empresa.auditoria import auditar_lectura
from empresa.permisos import requiere
from empresa.utils import contexto_selector, periodo_desde_request
from terceros.models import Tercero

from .models import ConceptoRetencion, Declaracion, DiferenciaFiscal, TarifaICA
from .reglas.formulario_350 import borrador_350 as calcular_350
from .reglas.ica import borrador_ica
from .reglas.renta import borrador_renta


class DeclaracionForm(forms.ModelForm):
    class Meta:
        model = Declaracion
        fields = ["tipo", "anio", "indice", "formulario", "valor_declarado", "estado", "elabora", "revisa", "firma", "presentada_el"]
        widgets = {"presentada_el": forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d")}


class ConceptoForm(forms.ModelForm):
    class Meta:
        model = ConceptoRetencion
        fields = ["codigo", "nombre", "base_minima_uvt", "tarifa_declarante", "tarifa_no_declarante",
                  "vigente_desde", "vigente_hasta", "estado", "fuente"]
        widgets = {"vigente_desde": forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d"),
                   "vigente_hasta": forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d")}


class DiferenciaForm(forms.ModelForm):
    class Meta:
        model = DiferenciaFiscal
        fields = ["anio", "concepto", "tipo", "valor", "soporte", "explicada"]


class TarifaICAForm(forms.ModelForm):
    class Meta:
        model = TarifaICA
        fields = ["ciiu", "descripcion", "tarifa_por_mil", "vigente_desde", "vigente_hasta", "estado", "fuente"]
        widgets = {"vigente_desde": forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d"),
                   "vigente_hasta": forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d")}


@requiere("ver_fiscal")
def fiscal(request):
    """Ruta anterior: la pantalla se separó en Renta, ICA, Exógena y Declaraciones."""
    return redirect(f"{reverse('renta')}?{request.GET.urlencode()}" if request.GET else reverse("renta"))


@requiere("ver_fiscal")
def renta(request):
    periodo = periodo_desde_request(request)
    r = borrador_renta(periodo.anio, periodo if periodo.saldos.exists() else None)
    ctx = {"renta": r.como_dict(), "diferencias": DiferenciaFiscal.objects.filter(anio=periodo.anio), "vencimientos": _vencimientos(["renta_c1", "renta_c2"], periodo.anio), "titulo": "Renta", **contexto_selector(periodo)}
    return render(request, "impuestos/renta.html", ctx)


@requiere("ver_fiscal")
def ica(request):
    periodo = periodo_desde_request(request)
    bimestre = (periodo.mes + 1) // 2
    ctx = {"ica": borrador_ica(periodo.anio, bimestre).como_dict(), "bimestre": bimestre, "vencimientos": _vencimientos(["ica"], periodo.anio), "titulo": "ICA", **contexto_selector(periodo)}
    return render(request, "impuestos/ica.html", ctx)


def _vencimientos(tipos, anio):
    """Fechas límite del año de los impuestos dados (el calendario es la fuente)."""
    from calendario.models import Obligacion

    return Obligacion.objects.filter(tipo__in=tipos, fecha_limite__year=anio).order_by("fecha_limite")


@requiere("ver_fiscal")
def reteica(request):
    periodo = periodo_desde_request(request)
    ctx = {"vencimientos": _vencimientos(["reteica"], periodo.anio), "titulo": "ReteICA", **contexto_selector(periodo)}
    return render(request, "impuestos/reteica.html", ctx)


def _obligacion_exogena(periodo, codigo):
    from analitica import anual

    _, total = anual.resultados_por_mes(periodo.anio)
    obligaciones, uvt, aviso_uvt = anual.obligacion_exogena(periodo.anio, total["ingresos"], solo=codigo)
    return obligaciones[0], total, uvt, aviso_uvt


@requiere("ver_fiscal")
def exogena_nacional(request):
    periodo = periodo_desde_request(request)
    obligacion, total, uvt, aviso_uvt = _obligacion_exogena(periodo, "EXOGENA_DIAN_UMBRAL_UVT")
    t = tope(periodo.anio)
    pagos = pagos_acumulados(periodo)
    maestro = {x.nit: x for x in Tercero.objects.filter(nit__in=list(pagos))}
    filas = []
    for nit, total_pagos in sorted(pagos.items(), key=lambda kv: -kv[1]):
        ter = maestro.get(nit)
        filas.append({"nit": nit, "nombre": ter.razon_social if ter else "", "total": total_pagos,
                      "reportable": (t is not None and total_pagos >= t),
                      "faltantes": ter.faltantes_exogena() if ter else ["no está en el maestro"]})
    ctx = {
        "obligacion": obligacion, "total": total, "uvt": uvt, "aviso_uvt": aviso_uvt, "tope_exogena": t, "exogena": filas[:100],
        "vencimientos": _vencimientos(["exogena"], periodo.anio), "titulo": "Exógena nacional (DIAN)", **contexto_selector(periodo),
    }
    return render(request, "impuestos/exogena_nacional.html", ctx)


@requiere("ver_fiscal")
def exogena_distrital(request):
    from analitica import anual

    periodo = periodo_desde_request(request)
    obligacion, total, uvt, aviso_uvt = _obligacion_exogena(periodo, "EXOGENA_DISTRITAL_UMBRAL_UVT")
    pagos = anual.pagos_por_tercero(periodo.anio)
    ctx = {
        "obligacion": obligacion, "total": total, "uvt": uvt, "aviso_uvt": aviso_uvt,
        "pagos": pagos[:150], "total_pagos": len(pagos), "sin_nit": sum(1 for f in pagos if not f["nit"]),
        "vencimientos": _vencimientos(["exogena_distrital"], periodo.anio), "titulo": "Exógena distrital (Bogotá)", **contexto_selector(periodo),
    }
    return render(request, "impuestos/exogena_distrital.html", ctx)


@requiere("ver_fiscal")
def declaraciones(request):
    periodo = periodo_desde_request(request)
    ctx = {"declaraciones": Declaracion.objects.filter(anio=periodo.anio), "titulo": "Declaraciones", **contexto_selector(periodo)}
    return render(request, "impuestos/declaraciones.html", ctx)


def _crud(request, modelo, formulario, titulo, destino, pk=None):
    obj = get_object_or_404(modelo, pk=pk) if pk else None
    form = formulario(request.POST or None, instance=obj)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Guardado.")
        return redirect(destino)
    return render(request, "empresa/formulario.html", {"form": form, "titulo": titulo})


@requiere("administrar")
def declaracion_editar(request, pk=None):
    return _crud(request, Declaracion, DeclaracionForm, "Declaración", "declaraciones", pk)


@requiere("administrar")
def concepto_editar(request, pk=None):
    return _crud(request, ConceptoRetencion, ConceptoForm, "Concepto de retención", "conceptos", pk)


@requiere("administrar")
def diferencia_editar(request, pk=None):
    return _crud(request, DiferenciaFiscal, DiferenciaForm, "Diferencia contable-fiscal", "renta", pk)


@requiere("administrar")
def tarifa_ica_editar(request, pk=None):
    return _crud(request, TarifaICA, TarifaICAForm, "Tarifa de ICA", "conceptos", pk)


@requiere("ver_fiscal")
def conceptos(request):
    return render(request, "impuestos/conceptos.html", {
        "conceptos": ConceptoRetencion.objects.all(), "tarifas_ica": TarifaICA.objects.all(),
        "titulo": "Conceptos de retención y tarifas de ICA",
    })


def _excel_350(datos, periodo):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Formulario 350"
    ws.append([f"BORRADOR · NO PRESENTAR · Formulario 350 · {periodo}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Casilla", "Concepto", "Personas", "Base", "Retención"])
    for celda in ws[2]:
        celda.font = Font(bold=True)
    for f in datos["filas"]:
        for lado, nombre in (("j", "Jurídicas"), ("n", "Naturales")):
            x = f[lado]
            if x and (x["base"] or x["retencion"]):
                ws.append([f"{x['casilla_base']} / {x['casilla_ret']}", f["etiqueta"], nombre, float(x["base"]), float(x["retencion"])])
    ws.append([])
    ws.append(["130", "Total retenciones renta y complementario", "", "", float(datos["total_renta"])])
    ws.append([str(datos["casilla_reteiva"]), "Retenciones de IVA (a responsables del IVA)", "", "", float(datos["reteiva"])])
    ws.append(["136", "Total retenciones", "", "", float(datos["total_retenciones"])])
    ws.append([])
    for aviso in datos["calculo"]["advertencias"] + datos["calculo"]["supuestos"]:
        ws.append([aviso])
    for col, ancho in zip("ABCDE", (14, 52, 12, 16, 16)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=3):
        for celda in fila[3:5]:
            if isinstance(celda.value, float):
                celda.number_format = "#,##0"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@requiere("ver_fiscal")
@auditar_lectura("Borrador del formulario 350")
def borrador_350(request):
    """Borrador del formulario 350 con los datos del balance y los auxiliares; se descarga en PDF o Excel."""
    from django.http import HttpResponse
    from django.template.loader import render_to_string

    from empresa.models import Empresa

    periodo = periodo_desde_request(request)
    datos = calcular_350(periodo)
    formato = request.GET.get("formato", "")
    if formato in ("pdf", "xlsx") and not request.user.puede("exportar"):
        from django.core.exceptions import PermissionDenied

        raise PermissionDenied
    ctx = {"datos": datos, "periodo": periodo, "empresa": Empresa.actual(), "titulo": "Borrador formulario 350", **contexto_selector(periodo)}
    if formato == "xlsx":
        contenido = _excel_350(datos, periodo)
        r = HttpResponse(contenido, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        r["Content-Disposition"] = f'attachment; filename="borrador-350-{periodo.anio}-{periodo.mes:02d}.xlsx"'
        return r
    if formato == "pdf":
        html = render_to_string("impuestos/borrador_350_pdf.html", ctx, request=request)
        try:
            from weasyprint import HTML

            r = HttpResponse(HTML(string=html).write_pdf(), content_type="application/pdf")
            r["Content-Disposition"] = f'attachment; filename="borrador-350-{periodo.anio}-{periodo.mes:02d}.pdf"'
            return r
        except Exception:  # noqa: BLE001 - sin WeasyPrint se entrega el HTML imprimible
            return HttpResponse(html)
    return render(request, "impuestos/borrador_350.html", ctx)
