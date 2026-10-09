"""Informe mensual: datos, Excel (openpyxl) y PDF (WeasyPrint si está disponible; si no, HTML imprimible)."""
import io
from decimal import Decimal

from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from analitica.servicios import indicadores_mes
from calendario.models import Obligacion
from controles.models import Hallazgo
from empresa.models import Empresa


def datos_informe(periodo):
    hallazgos = Hallazgo.objects.filter(periodo=periodo).select_related("regla")
    abiertos = hallazgos.filter(estado="abierto")
    proximas = (
        Obligacion.objects.filter(fecha_limite__gte=timezone.localdate())
        .exclude(estado__in=["presentada", "pagada"]).order_by("fecha_limite")[:8]
    )
    return {
        "empresa": Empresa.actual(), "periodo": periodo, "generado": timezone.localtime(),
        "indicadores": indicadores_mes(periodo),
        "hallazgos_abiertos": abiertos.order_by("severidad"),
        "conteo": {
            "alta": abiertos.filter(severidad="alta").count(), "media": abiertos.filter(severidad="media").count(),
            "baja": abiertos.filter(severidad="baja").count(),
            "explicados": hallazgos.filter(estado="explicado").count(), "corregidos": hallazgos.filter(estado="corregido").count(),
        },
        "proximas": proximas,
    }


def semaforo_general(conteo):
    if conteo["alta"]:
        return "rojo"
    return "ambar" if conteo["media"] else "verde"


def generar_pdf(periodo):
    """Devuelve (bytes, content_type). Si WeasyPrint no está disponible, entrega HTML imprimible."""
    ctx = datos_informe(periodo)
    ctx["semaforo"] = semaforo_general(ctx["conteo"])
    html = render_to_string("reportes/informe.html", ctx)
    try:
        from weasyprint import HTML

        return HTML(string=html).write_pdf(), "application/pdf"
    except Exception:
        return html.encode("utf-8"), "text/html; charset=utf-8"


def _fmt(v):
    return float(v) if isinstance(v, Decimal) else v


def generar_excel(periodo):
    ctx = datos_informe(periodo)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    negrita = Font(bold=True)
    ws.append([f"Informe mensual {periodo} · {ctx['empresa'] or ''}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Generado", ctx["generado"].strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(["Indicador", "Actual", "Mismo mes año anterior"])
    for c in ws[4]:
        c.font = negrita
        c.fill = PatternFill("solid", fgColor="DCE3F0")
    ind = ctx["indicadores"]
    for clave in sorted((ind.get("actual") or {}).keys()):
        ws.append([clave.replace("_", " "), _fmt(ind["actual"].get(clave)), _fmt((ind.get("anio_anterior") or {}).get(clave))])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 24

    h = wb.create_sheet("Hallazgos")
    h.append(["Regla", "Severidad", "Estado", "Título", "Detalle", "Explicación"])
    for c in h[1]:
        c.font = negrita
    for x in Hallazgo.objects.filter(periodo=periodo).select_related("regla"):
        h.append([x.regla.codigo, x.get_severidad_display(), x.get_estado_display(), x.titulo, x.detalle, x.explicacion])
    for col, w in zip("ABCDEF", (10, 14, 12, 50, 70, 50)):
        h.column_dimensions[col].width = w
    for fila in h.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    v = wb.create_sheet("Vencimientos")
    v.append(["Obligación", "Período", "Fecha límite", "Estado", "Verificación", "Elabora", "Revisa", "Firma"])
    for c in v[1]:
        c.font = negrita
    for o in ctx["proximas"]:
        v.append([o.nombre, o.periodo_texto, o.fecha_limite, o.get_estado_display(), o.get_verificacion_display(), o.elabora, o.revisa, o.firma])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def tabla_a_excel(titulo, encabezados, filas):
    """Exportación genérica de una tabla de la interfaz a Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append([_fmt(x) for x in f])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
