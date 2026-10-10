import io
from datetime import date
from decimal import Decimal

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from asistente.servicio import explicacion_local, explicar, minimizar
from conciliaciones import servicios
from conciliaciones.models import MovimientoBanco
from contabilidad.models import Movimiento
from controles.models import Hallazgo, ReglaControl
from empresa.models import Empresa, Parametro
from reportes.informe import generar_excel, generar_pdf

from .helpers import archivo, auxiliar, balance, factura, periodo


@pytest.mark.django_db
def test_conciliar_banco(datos_iniciales):
    p = periodo()
    auxiliar(p, [
        (date(2026, 9, 1), "R1", "", "111005", "", 1000, 0),
        (date(2026, 9, 5), "E1", "", "111005", "", 0, 300),
        (date(2026, 9, 9), "E2", "", "111005", "", 0, 50),
    ])
    a = archivo("extracto_banco", p)
    for f, v in ((date(2026, 9, 2), 1000), (date(2026, 9, 5), -300), (date(2026, 9, 20), -77)):
        MovimientoBanco.objects.create(periodo=p, archivo=a, fecha=f, valor=Decimal(v))
    r = servicios.conciliar_banco(p, "1110")
    assert r["conciliadas"] == 2 and len(r["solo_banco"]) == 1 and len(r["solo_libros"]) == 1


@pytest.mark.django_db
def test_conciliar_facturas_y_terceros(datos_iniciales):
    p = periodo()
    factura(p, "recibida", "1", date(2026, 9, 5), "900111", "900902549", 1000, 190, prefijo="F")
    factura(p, "recibida", "2", date(2026, 9, 6), "900111", "900902549", 1000, 190, prefijo="F")
    auxiliar(p, [(date(2026, 9, 5), "C1", "F1", "220505", "900111", 0, 1190), (date(2026, 9, 5), "C1", "F1", "513595", "900111", 1190, 0)])
    r = servicios.conciliar_facturas(p)
    compras = r["filas"][1]
    assert compras["facturas"] == 2 and compras["causadas"] == 1 and compras["sin_causar"] == 1
    t = servicios.conciliar_terceros(p, ["22"], -1)
    assert t["total"] == Decimal("1190")
    assert servicios.conciliar_auxiliares_vs_balance(p)["filas"]


@pytest.mark.django_db
def test_informe_excel_y_pdf(datos_iniciales):
    p = periodo()
    balance(p, [("1105", "Caja", 0, 0, 0, -10)])
    from controles.motor import ejecutar_reglas

    ejecutar_reglas(p)
    wb = load_workbook(io.BytesIO(generar_excel(p)))
    assert wb.sheetnames == ["Resumen", "Hallazgos", "Vencimientos"]
    assert wb["Hallazgos"].max_row > 1
    contenido, tipo = generar_pdf(p)
    assert contenido and tipo in ("application/pdf", "text/html; charset=utf-8")
    if tipo == "application/pdf":
        assert contenido.startswith(b"%PDF")


@pytest.mark.django_db
def test_descarga_informe_por_http(cliente_dueno, datos_iniciales):
    r = cliente_dueno.get("/informes/excel/?anio=2026&mes=9")
    assert r.status_code == 200 and r["Content-Type"].endswith("sheet")
    r = cliente_dueno.get("/informes/pdf/?anio=2026&mes=9")
    assert r.status_code == 200


def test_minimizar_quita_identificadores():
    t = minimizar("NIT 900.902.549-7, cédula 1020304050, correo juan@empresa.com, valor 1.500")
    assert "900.902.549" not in t and "1020304050" not in t and "juan@" not in t and "1.500" in t


@pytest.mark.django_db
def test_asistente_sin_clave_usa_plantilla(datos_iniciales, settings):
    from controles.motor import sincronizar_catalogo

    sincronizar_catalogo()
    h = Hallazgo.objects.create(regla=ReglaControl.objects.get(codigo="INT001"), periodo=periodo(), clave="x",
                                titulo="Débitos distintos de créditos", detalle="Diferencia 10", severidad="alta")
    settings.ANTHROPIC_API_KEY = ""
    e = explicar(h)
    assert e.modelo == "local" and "Diferencia 10" in e.texto
    assert "contador" in explicacion_local(h)


@pytest.mark.django_db
def test_asistente_con_api_envia_solo_datos_minimos(datos_iniciales, settings, monkeypatch):
    from asistente import servicio
    from controles.motor import sincronizar_catalogo

    sincronizar_catalogo()
    h = Hallazgo.objects.create(regla=ReglaControl.objects.get(codigo="TER001"), periodo=periodo(), clave="x",
                                titulo="DV errado para 900.902.549-7", detalle="juan@x.com", severidad="media")
    settings.ANTHROPIC_API_KEY = "clave-de-prueba"
    enviado = {}

    def falso(payload):
        enviado.update(payload)
        return {"content": [{"type": "text", "text": "Explicación redactada."}]}

    monkeypatch.setattr(servicio, "_llamar_api", falso)
    e = explicar(h)
    assert e.texto == "Explicación redactada."
    cuerpo = enviado["messages"][0]["content"]
    assert "900.902.549" not in cuerpo and "juan@x.com" not in cuerpo
    assert "NO calcules" in enviado["system"]


@pytest.mark.django_db
def test_flujo_http_de_carga_y_hallazgos(cliente_dueno, datos_iniciales):
    csv = b"cuenta,debito,credito,saldo_final\n1105,100,0,100\n2205,0,90,90\n"
    r = cliente_dueno.post("/cargas/nueva/", {"tipo": "balance", "anio": 2026, "mes": 9, "archivo": SimpleUploadedFile("b.csv", csv)})
    assert r.status_code == 302
    detalle = cliente_dueno.get(r.url)
    assert detalle.status_code == 200 and b"Vista previa" in detalle.content
    r2 = cliente_dueno.post(r.url + "confirmar/")
    assert r2.status_code == 302
    h = cliente_dueno.get("/hallazgos/?anio=2026&mes=9")
    assert b"d\xc3\xa9bitos no son iguales" in h.content.lower() or b"INT001" in h.content
    hallazgo = Hallazgo.objects.filter(regla__codigo="INT001").first()
    r3 = cliente_dueno.post(f"/hallazgos/{hallazgo.pk}/", {"accion": "estado", "estado": "explicado", "explicacion": "Ajuste pendiente de Ideako"})
    hallazgo.refresh_from_db()
    assert hallazgo.estado == "explicado" and r3.status_code == 302
    sin = cliente_dueno.post(f"/hallazgos/{hallazgo.pk}/", {"accion": "estado", "estado": "explicado", "explicacion": ""})
    assert sin.status_code == 302


@pytest.mark.django_db
def test_tablero_con_datos(cliente_dueno, datos_iniciales):
    p = periodo()
    balance(p, [("4135", "Ventas", 0, 0, 500, 500), ("1105", "Caja", 0, 500, 0, 500)])
    r = cliente_dueno.get("/?anio=2026&mes=9")
    assert r.status_code == 200 and b"Ingresos acumulados" in r.content
    assert Empresa.actual().nit_formateado == "900.902.549-7"
    Parametro.objects.filter(codigo="RENTA_TARIFA").delete()
    assert cliente_dueno.get("/renta/?anio=2026&mes=9").status_code == 200


def test_excel_neutraliza_formulas():
    from reportes.informe import tabla_a_excel

    wb = load_workbook(io.BytesIO(tabla_a_excel("T", ["a"], [("=HYPERLINK(\"http://x\")",), ("normal",), (Decimal("5"),)])))
    ws = wb.active
    assert ws["A2"].value.startswith("'=") and ws["A3"].value == "normal" and ws["A4"].value == 5


@pytest.mark.django_db
def test_conciliar_banco_fecha_lejana_gastos_agrupados_y_sumas(datos_iniciales):
    p = periodo()
    auxiliar(p, [
        (date(2026, 9, 28), "R1", "", "111005", "", 500, 0),            # el banco lo movió el 2/9: misma cifra, otra fecha
        (date(2026, 9, 30), "G1", "", "111005", "", 0, "19.00"),        # un solo asiento de gravamen y comisiones
        (date(2026, 9, 10), "E1", "", "111005", "", 0, 700),            # el banco paga 1000 y en libros se separó un descuento
        (date(2026, 9, 10), "E2", "", "111005", "", 0, 300),
        (date(2026, 9, 12), "E3", "", "111005", "", 0, "40.50"),        # diferencia de centavos con el extracto
    ])
    Movimiento.objects.filter(comprobante="G1").update(descripcion="GRAVAMEN FINANCIERO Y COMISIONES")
    a = archivo("extracto_banco", p)
    banco = [
        (date(2026, 9, 2), 500, "TRANSFERENCIA"), (date(2026, 9, 3), "-10.00", "IMPTO GOBIERNO 4X1000"),
        (date(2026, 9, 9), "-9.00", "COBRO IVA PAGOS AUTOMATICOS"), (date(2026, 9, 10), -1000, "PAGO A PROVEEDOR"),
        (date(2026, 9, 12), "-40.00", "PAGO PSE"), (date(2026, 9, 20), -77, "OTRO"),
    ]
    for f, v, d in banco:
        MovimientoBanco.objects.create(periodo=p, archivo=a, fecha=f, valor=Decimal(str(v)), descripcion=d)
    r = servicios.conciliar_banco(p, "1110")
    motivos = [g["motivo"] for g in r["con_revision"]]
    assert r["exactas"] == 1 and len(r["con_revision"]) == 3  # el pago con centavos de diferencia queda exacto
    assert any("días de diferencia" in m for m in motivos) and any("Gastos bancarios" in m for m in motivos)
    assert any("suma de 2 partidas de libros" in m for m in motivos)
    assert [b.valor for b in r["solo_banco"]] == [Decimal("-77")] and r["solo_libros"] == []


@pytest.mark.django_db
def test_factura_causada_con_el_consecutivo_de_world_office(datos_iniciales):
    """World Office registra «(DTS) FV FE 11407» y «(DTS) FC DHT 1315», no el número de la factura de la DIAN."""
    from controles.reglas_facturas import indice_documentos, movimientos_de, valor_causado

    p = periodo()
    venta = factura(p, "emitida", "11407", date(2026, 9, 11), "900902549", "800111222", 378067, 71833, prefijo="FE")
    compra = factura(p, "recibida", "EB1", date(2026, 9, 7), "899999115", "900902549", 97680, 18559, nombre_emisor="EMPRESA DE TELECOMUNICACIONES DE BOGOTA SA ESP SIGLA ETB")
    auxiliar(p, [
        (date(2026, 9, 11), "FV", "(DTS) FV FE 11407", "13050501", "", 449900, 0),
        (date(2026, 9, 11), "FV", "(DTS) FV FE 11407", "13551511", "", "2079.37", 0),      # anticipo de retención: no es cartera
        (date(2026, 9, 11), "FV", "(DTS) FV FE 11407", "41355402", "", 0, 378067),
        (date(2026, 9, 11), "RC", "(DTS) RC DH 7411", "13050501", "", 0, 449900),           # el recibo de caja baja la cartera
        (date(2026, 9, 11), "FC", "(DTS) FC DHT 1315", "23355001", "", 0, "116239.62"),
        (date(2026, 9, 11), "FC", "(DTS) FC DHT 1315", "51353501", "", "97680.35", 0),
    ])
    Movimiento.objects.filter(documento="(DTS) FC DHT 1315").update(tercero_nombre="EMPRESA DE TELECOMUNICACIONES ETB")
    indice = indice_documentos(p)
    assert movimientos_de(venta, indice)                              # por el número dentro del documento
    assert valor_causado(venta, movimientos_de(venta, indice)) == Decimal("449900")
    assert movimientos_de(compra, indice)                             # por tercero y valor (±$5)
    sin = factura(p, "recibida", "X9", date(2026, 9, 8), "899999999", "900902549", 1000, 190, nombre_emisor="OTRO PROVEEDOR SAS")
    assert not movimientos_de(sin, indice)


@pytest.mark.django_db
def test_libro_auxiliar_de_todo_el_anio_se_reparte_por_mes_y_completa_la_exogena(datos_iniciales, cliente_dueno):
    from datetime import date as d

    from controles.reglas_exogena import anio_completo, meses_con_auxiliares
    from empresa import cargas
    from empresa.models import ArchivoCargado, Periodo

    p12 = Periodo.obtener(2025, 12)
    a = ArchivoCargado.objects.create(tipo="auxiliar", nombre_original="aux2025.xlsx", hash_sha256="a" * 64, tamano=1, periodo=p12,
                                      varios_meses=True, estado="pendiente", archivo="cargas/x/a.xlsx")
    filas = [{"fecha": d(2025, m, 15), "cuenta": "513535", "debito": Decimal("100"), "credito": Decimal("0")} for m in range(1, 13)]
    filas.append({"fecha": d(2025, 3, 2), "cuenta": "413554", "debito": Decimal("0"), "credito": Decimal("5000")})
    from contabilidad.importadores import importar_auxiliar

    resumen = importar_auxiliar(a, filas)
    a.estado, a.resumen, a.vigente = "importado", resumen, True
    a.save()
    assert resumen["meses"][0] == "2025-01" and len(resumen["meses"]) == 12
    assert meses_con_auxiliares(2025) == set(range(1, 13)) and anio_completo(p12)
    assert Movimiento.objects.filter(periodo__mes=3, periodo__anio=2025).count() == 2

    html = cliente_dueno.get("/anual/?anio=2025").content.decode()
    assert "12 de 12 meses" in html and "Exógena nacional (DIAN)" in html and "Exógena distrital (Bogotá)" in html
    assert "Pendiente" in html  # DIAN: falta el umbral; distrital: 3.500 UVT × UVT 2025 no se alcanza con $5.000
    assert "No alcanza el umbral" in html

    mensual = ArchivoCargado.objects.create(tipo="auxiliar", nombre_original="m.xlsx", hash_sha256="b" * 64, tamano=1,
                                            periodo=Periodo.obtener(2025, 6), estado="pendiente", archivo="cargas/x/m.xlsx")
    with pytest.raises(cargas.ConflictoDeMeses):
        cargas._verificar_conflicto_de_meses_auxiliar(mensual, [])
